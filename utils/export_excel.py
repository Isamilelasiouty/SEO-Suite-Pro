"""
Excel Export Utility for SEO Suite Pro:
Generates multi-sheet, gold-accented .xlsx workbooks using openpyxl with auto-fitted column widths and custom typography.
"""

import io
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def generate_excel_bytes(sheets_dict: dict[str, pd.DataFrame]) -> bytes:
    """
    Creates a styled Excel workbook from a dictionary of {sheet_name: DataFrame}.
    Styling adheres to the executive gold and dark theme.
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, df in sheets_dict.items():
            clean_name = sheet_name[:31]  # Excel limits sheet names to 31 chars
            df_to_write = df if not df.empty else pd.DataFrame([{"Notice": "No data available"}])
            df_to_write.to_excel(writer, sheet_name=clean_name, index=False)
            
            worksheet = writer.sheets[clean_name]
            
            # Gold header styling
            header_fill = PatternFill(start_color="D4AF37", end_color="D4AF37", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="000000")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            thin_border = Border(
                left=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"),
                bottom=Side(style="thin", color="CCCCCC")
            )
            
            # Format header row
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align
                cell.border = thin_border
                
            # Format data rows
            body_font = Font(name="Segoe UI", size=10, color="222222")
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                for cell in row:
                    cell.font = body_font
                    cell.border = thin_border
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00"
                    elif isinstance(cell.value, int):
                        cell.number_format = "#,##0"
                        
            # Auto-adjust column widths
            for col in worksheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                worksheet.column_dimensions[col_letter].width = max(12, min(50, max_len + 3))
                
    output.seek(0)
    return output.getvalue()
